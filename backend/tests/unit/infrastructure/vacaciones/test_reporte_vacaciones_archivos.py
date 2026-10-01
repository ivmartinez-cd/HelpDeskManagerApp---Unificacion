"""Excel y PDF del reporte de vacaciones: se generan y el Excel no ejecuta fórmulas."""

import io

from openpyxl import load_workbook

from src.modules.vacaciones.application.dtos.reporte_dtos import (
    FilaEmpleadoReporteDTO,
    FilaSectorReporteDTO,
    ReporteVacacionesDTO,
)
from src.modules.vacaciones.infrastructure.exports.reporte_vacaciones_archivos import (
    reporte_excel,
    reporte_pdf,
)


def _reporte() -> ReporteVacacionesDTO:
    empleado = FilaEmpleadoReporteDTO(
        nombre="=HYPERLINK(\"x\")", color="#000", sector_nombre="Mesa", cargo_nombre="TL",
        annual=14, used=3, pending=1, available=10,
    )
    sector = FilaSectorReporteDTO(
        nombre="Mesa", color="#000", empleados=1, annual=14, used=3, available=10
    )
    return ReporteVacacionesDTO(year=2026, por_empleado=[empleado], por_sector=[sector])


def test_excel_tiene_las_dos_hojas_y_no_ejecuta_formulas() -> None:
    libro = load_workbook(io.BytesIO(reporte_excel(_reporte())))
    assert libro.sheetnames == ["Por empleado", "Por sector"]
    assert libro["Por empleado"]["A2"].value == "'=HYPERLINK(\"x\")"
    assert libro["Por empleado"]["G2"].value == 10


def test_pdf_se_genera() -> None:
    assert reporte_pdf(_reporte(), timezone="America/Argentina/Buenos_Aires").startswith(b"%PDF")
