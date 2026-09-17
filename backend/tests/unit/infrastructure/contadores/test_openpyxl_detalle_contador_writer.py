"""OpenpyxlDetalleContadorWriter: título con cliente/proceso, header en
negrita en fila 3, 15 columnas en el mismo orden de la pantalla, una fila
por `DetalleContadorRow`, y semaforización roja en "Tipo" cuando falta
contador."""

import io

from openpyxl import load_workbook

from src.modules.contadores.domain.value_objects.detalle_contador_row import (
    DetalleContadorRow,
)
from src.modules.contadores.infrastructure.xlsx.openpyxl_detalle_contador_writer import (
    OpenpyxlDetalleContadorWriter,
)


def _fila(falta_contador: bool = True) -> DetalleContadorRow:
    return DetalleContadorRow(
        empresa="ISSN",
        sucursal="Botiquin Caviahue",
        sector="Indeterminado",
        modelo="MFP Mono HP 432fdn",
        serie="CNB1R4C0MV",
        nombre_clase="Mono",
        fecha_toma_anterior=None,
        contador_anterior=1,
        fecha_toma_actual=None,
        contador_actual=0,
        impresiones_reales=0.0,
        estado_maquina="Activa en Cliente",
        direccion_ip=None,
        mascara_ip=None,
        falta_contador=falta_contador,
        tipo="FALTA CONTADOR Mono" if falta_contador else None,
        nro_proceso=99089,
        nombre_anexo="Anexo Principal",
        periodo_facturacion="2026-08",
    )


def test_write_arma_titulo_header_y_una_fila_por_registro() -> None:
    contenido = OpenpyxlDetalleContadorWriter().write([_fila(), _fila()], "ISSN", 99089)

    wb = load_workbook(io.BytesIO(contenido))
    ws = wb.active
    assert ws is not None
    assert ws["A1"].value == "Detalle de Contadores — ISSN · Proceso 99089"

    header = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    assert header == [
        "Empresa",
        "Sucursal",
        "Modelo",
        "Serie",
        "Sector",
        "Fecha Toma Ant.",
        "Contador Ant.",
        "Fecha Toma Act.",
        "Contador Act.",
        "Impresiones",
        "Tipo",
        "Clase",
        "Estado Máquina",
        "Dirección IP",
        "Máscara IP",
    ]
    assert next(ws.iter_rows(min_row=3, max_row=3))[0].font.bold is True
    assert ws.max_row == 5
    primera_fila = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert primera_fila[0] == "ISSN"
    assert primera_fila[3] == "CNB1R4C0MV"
    assert primera_fila[10] == "FALTA CONTADOR Mono"


def test_write_marca_en_rojo_la_columna_tipo_cuando_falta_contador() -> None:
    contenido = OpenpyxlDetalleContadorWriter().write(
        [_fila(falta_contador=True), _fila(falta_contador=False)], "ISSN", 99089
    )

    wb = load_workbook(io.BytesIO(contenido))
    ws = wb.active
    assert ws is not None
    celda_falta = ws.cell(row=4, column=11)
    celda_ok = ws.cell(row=5, column=11)
    assert celda_falta.font.color.rgb == "00EF4444"
    assert celda_falta.font.bold is True
    assert celda_ok.font.bold is not True


def test_write_sin_filas_deja_solo_titulo_y_header() -> None:
    contenido = OpenpyxlDetalleContadorWriter().write([], "ISSN", 99089)

    wb = load_workbook(io.BytesIO(contenido))
    ws = wb.active
    assert ws is not None
    assert ws.max_row == 3
