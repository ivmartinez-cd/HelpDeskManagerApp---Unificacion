"""Escritor del reporte ejecutivo XLSX de "Detalle de contadores por
proceso" — mismas 15 columnas que la pantalla (pedido explícito del
usuario, sin resumir para el cliente), armado en memoria: se descarga en la
misma request, a diferencia de las herramientas de `tools_router.py` que sí
necesitan un archivo intermedio en disco.

Formato ejecutivo con la línea Institucional de Canal Directo (naranja
`#F7941D` + gris MPS `#58595B`, ver "Manual de marca.pdf") — la misma
paleta que usa la pantalla, más semaforización roja para "Falta Contador"
(excepción de marca ya aceptada para semáforos de estado, ver
`feedback_brand_purity_canal_directo`)."""

import io
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet

from src.modules.contadores.domain.value_objects.detalle_contador_row import (
    DetalleContadorRow,
)

_NARANJA = "F7941D"
_GRIS = "58595B"
_ROJO_FALTA = "EF4444"
_GRIS_ZEBRA = "F2F2F2"
_BORDE = "DDDDDD"
_ISOTIPO = Path(__file__).parent / "assets" / "isotipo-white.png"
_ARGENTINA_TZ = ZoneInfo("America/Argentina/Buenos_Aires")

_COLUMNAS: tuple[tuple[str, int], ...] = (
    ("Empresa", 22),
    ("Sucursal", 22),
    ("Modelo", 26),
    ("Serie", 18),
    ("Sector", 20),
    ("Fecha Toma Ant.", 15),
    ("Contador Ant.", 14),
    ("Fecha Toma Act.", 15),
    ("Contador Act.", 14),
    ("Impresiones", 13),
    ("Tipo", 20),
    ("Clase", 10),
    ("Estado Máquina", 18),
    ("Dirección IP", 15),
    ("Máscara IP", 15),
)
_NUMERICAS = {"Contador Ant.", "Contador Act.", "Impresiones"}
_FECHAS = {"Fecha Toma Ant.", "Fecha Toma Act."}
_ULTIMA_COL = get_column_letter(len(_COLUMNAS))
_FILA_HEADER = 3


class OpenpyxlDetalleContadorWriter:
    def write(self, filas: list[DetalleContadorRow], cliente: str, nro_proceso: int) -> bytes:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Detalle de contadores"
        _titulo(ws, cliente, nro_proceso, len(filas))
        _encabezado(ws)
        for i, fila in enumerate(filas):
            _escribir_fila(ws, _FILA_HEADER + 1 + i, fila)
        ws.freeze_panes = f"A{_FILA_HEADER + 1}"
        ws.auto_filter.ref = f"A{_FILA_HEADER}:{_ULTIMA_COL}{_FILA_HEADER}"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _titulo(ws: Worksheet, cliente: str, nro_proceso: int, total: int) -> None:
    ws.merge_cells(f"A1:{_ULTIMA_COL}1")
    ws["A1"] = f"Detalle de Contadores — {cliente} · Proceso {nro_proceso}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=_NARANJA)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=4)
    ws.row_dimensions[1].height = 26
    _isotipo(ws)

    ws.merge_cells(f"A2:{_ULTIMA_COL}2")
    generado = datetime.now(_ARGENTINA_TZ).strftime("%d/%m/%Y %H:%M")
    ws["A2"] = f"Generado el {generado} · {total} equipo{'s' if total != 1 else ''}"
    ws["A2"].font = Font(italic=True, size=10, color=_GRIS)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")


def _isotipo(ws: Worksheet) -> None:
    """Isotipo blanco de Canal Directo (mismo asset que el header del
    sidebar, `frontend/public/isotipo-white.svg`, rasterizado a PNG — ver
    `assets/isotipo-white.png`) sobre la barra de título naranja."""
    img = XLImage(str(_ISOTIPO))
    lado = pixels_to_EMU(18)
    marca = AnchorMarker(col=0, row=0, colOff=pixels_to_EMU(6), rowOff=pixels_to_EMU(8))
    img.anchor = OneCellAnchor(_from=marca, ext=XDRPositiveSize2D(lado, lado))
    ws.add_image(img)


def _encabezado(ws: Worksheet) -> None:
    ws.row_dimensions[_FILA_HEADER].height = 30
    for idx, (nombre, ancho) in enumerate(_COLUMNAS):
        col = get_column_letter(idx + 1)
        ws.column_dimensions[col].width = ancho
        celda = ws[f"{col}{_FILA_HEADER}"]
        celda.value = nombre
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_GRIS)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _escribir_fila(ws: Worksheet, num_fila: int, f: DetalleContadorRow) -> None:
    zebra = PatternFill("solid", fgColor=_GRIS_ZEBRA) if num_fila % 2 == 0 else None
    borde = Border(bottom=Side(style="thin", color=_BORDE))
    for idx, ((nombre, _), valor) in enumerate(zip(_COLUMNAS, _fila_a_excel(f), strict=True)):
        celda = ws.cell(row=num_fila, column=idx + 1, value=valor)
        celda.border = borde
        if zebra:
            celda.fill = zebra
        if nombre in _NUMERICAS:
            celda.alignment = Alignment(horizontal="right")
        elif nombre in _FECHAS:
            celda.alignment = Alignment(horizontal="center")
        if nombre == "Tipo" and f.falta_contador:
            celda.font = Font(bold=True, color=_ROJO_FALTA)


def _fila_a_excel(f: DetalleContadorRow) -> list[object]:
    return [
        f.empresa,
        f.sucursal,
        f.modelo,
        f.serie,
        f.sector or "",
        _fecha(f.fecha_toma_anterior),
        f.contador_anterior,
        _fecha(f.fecha_toma_actual),
        f.contador_actual,
        f.impresiones_reales,
        f.tipo or "",
        f.nombre_clase or "",
        f.estado_maquina or "",
        f.direccion_ip or "",
        f.mascara_ip or "",
    ]


def _fecha(valor: date | None) -> str:
    return valor.strftime("%d/%m/%Y") if valor else ""
